"""Export a run to Excel, one row per diagnosis.
    python3 export_run_excel.py results/<run>.json [results/<run>.xlsx]
Needs openpyxl (pip install openpyxl). Candidates are those recorded in the run."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from icd_tree import Tree
T = Tree(); d = lambda c: T.node[c]['desc'] if c in T.node else ''
run = sys.argv[1]; out = sys.argv[2] if len(sys.argv) > 2 else run.replace('.json', '.xlsx')
U = json.load(open(run))['units']
wb = Workbook(); ws = wb.active; ws.title = 'Diagnoses'
cols = ['diagnosis id', 'chart (med_id)', 'phrase', 'facts sent', 'system chosen', 'category chosen', 'category description',
        'entry reached', 'entry description', 'candidates offered (code: source)', 'final code', 'final description',
        'coder code', 'coder description', 'engine code', 'coder kept the code', 'coder category in candidates', 'result']
ws.append(cols)
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4FD6')
for u in U:
    w = u['runs'][0]; p = {x['level']: x for x in w['path']}
    entry = next((p[l]['resolved'] for l in ('flat', 'sub3', 'sub2', 'sub1', 'category') if l in p and p[l].get('resolved')), None)
    cand = p.get('candidates'); cs = cand['resolved'] if cand else ([entry] if entry else [])
    origins = cand['answer']['origins'] if cand else ['code-based walk'] * len(cs)
    origins = [o.replace('facility link', 'custom link').replace('tabular note', 'Excludes1/2 note') for o in origins]
    g = u.get('coder_code'); cat = (p.get('category') or {}).get('resolved')
    covered = bool(g) and g[:3] in {c[:3] for c in cs}
    if w.get('error'): res = 'API error'
    elif not entry: res = 'no entry - Gemini route'
    elif not u.get('coder_kept'): res = 'coder removed the code - not scored'
    else: res = 'covered' if covered else 'MISSED'
    ws.append([u['uid'], u['med_id'], u['phrase'], json.dumps({k: v for k, v in u['state']['diagnosis'].items() if k != 'phrase'}, ensure_ascii=False),
               (p.get('system') or {}).get('resolved'), cat, d(cat), entry, d(entry), '; '.join(f'{c}: {o}' for c, o in zip(cs, origins)),
               w.get('code'), d(w.get('code')), g, d(g), u.get('engine_code'), 'yes' if u.get('coder_kept') else 'no', 'yes' if covered else 'no', res])
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
for i, wd in enumerate([16, 18, 40, 40, 10, 10, 34, 10, 34, 70, 10, 34, 10, 34, 10, 8, 8, 22], 1): ws.column_dimensions[get_column_letter(i)].width = wd
wb.save(out); print('saved', out, len(U), 'rows')
